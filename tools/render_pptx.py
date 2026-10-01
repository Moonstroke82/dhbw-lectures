"""Render a laid-out Deck to an editable .pptx (python-pptx)."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt
from lxml import etree

import theme
from slidemd import inline_runs

C = {k: RGBColor.from_string(v) for k, v in theme.COLORS.items()}


def _box(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tb, tf


def _runs(p, text, size, color="text", bold=False, italic=False):
    for t, st in inline_runs(text):
        r = p.add_run()
        r.text = t
        f = r.font
        f.size = Pt(size)
        f.name = theme.FONT_MONO if st.get("code") else theme.FONT
        f.bold = bold or st.get("b", False)
        f.italic = italic or st.get("i", False)
        f.color.rgb = C["accent"] if st.get("link") else C[color]
        if st.get("link"):
            r.hyperlink.address = st["link"]


def _bullet(p, level, ordered, hanging_only=False):
    pPr = p._p.get_or_add_pPr()
    ind = Pt(theme.BULLET_INDENT * (level + 1))
    pPr.set("marL", str(int(ind)))
    pPr.set("indent", str(-int(Pt(theme.BULLET_INDENT))))
    if hanging_only:
        etree.SubElement(pPr, qn("a:buNone"))
        return
    clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr")).set("val", theme.COLORS["accent"])
    if ordered:
        etree.SubElement(pPr, qn("a:buFont")).set("typeface", "+mj-lt")
        etree.SubElement(pPr, qn("a:buAutoNum")).set("type", "arabicPeriod")
    else:
        etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
        etree.SubElement(pPr, qn("a:buChar")).set("char", "•" if level == 0 else "–")


def _text(slide, el, refs=False):
    _, tf = _box(slide, el["x"], el["y"], el["w"], el["h"] + 4)
    color = "muted" if el.get("muted") else "text"
    for i, para in enumerate(el["paras"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        size = para.get("size", el["size"])
        if "space" in para:
            p.space_after = Pt(para["space"])
        if "level" in para:
            _bullet(p, para["level"], para["ordered"], hanging_only=refs)
        _runs(p, para["text"], size, color=para.get("color", color), bold=para.get("bold", False))


def _rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Pt(x), Pt(y), Pt(w), Pt(h))
    s.fill.solid()
    s.fill.fore_color.rgb = C[fill]
    if line:
        s.line.color.rgb = C[line]
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.08, 8 / max(min(w, h), 1))
    return s


def _table(slide, el):
    b = el["block"]
    rows = [b["header"]] + b["rows"]
    shape = slide.shapes.add_table(len(rows), len(b["header"]), Pt(el["x"]), Pt(el["y"]), Pt(el["w"]), Pt(el["h"]))
    tbl = shape.table
    for c, wdt in enumerate(el["colw"]):
        tbl.columns[c].width = Pt(wdt)
    for r, rh in enumerate(el["rowh"]):
        tbl.rows[r].height = Pt(rh)
        for c in range(len(b["header"])):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = cell.margin_top = cell.margin_bottom = Pt(theme.CELL_PAD)
            cell.fill.solid()
            cell.fill.fore_color.rgb = C["primary"] if r == 0 else (C["light"] if r % 2 == 0 else C["white"])
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[b["aligns"][c]]
            txt = rows[r][c] if c < len(rows[r]) else ""
            _runs(p, txt, el["size"], color="white" if r == 0 else "text", bold=(r == 0))


def _code(slide, el):
    s = _rect(slide, el["x"], el["y"], el["w"], el["h"], "light", "border")
    tf = s.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Pt(theme.BOX_PAD)
    for i, ln in enumerate(el["text"].splitlines() or [""]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = ln
        r.font.name = theme.FONT_MONO
        r.font.size = Pt(el["size"])
        r.font.color.rgb = C["text"]


def _layer(slide, el):
    fill = "primary" if el["index"] % 2 == 0 else "accent"
    _rect(slide, el["x"], el["y"], el["lw"], el["h"], fill)
    _rect(slide, el["x"] + el["lw"], el["y"], el["w"] - el["lw"], el["h"], "light")
    pad = theme.BOX_PAD
    _, tf = _box(slide, el["x"] + pad, el["y"] + pad, el["lw"] - 2 * pad, el["h"] - 2 * pad)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _runs(tf.paragraphs[0], el["label"], el["size"], color="white", bold=True)
    if el["desc"]:
        _, tf = _box(slide, el["x"] + el["lw"] + pad, el["y"] + pad, el["w"] - el["lw"] - 2 * pad, el["h"] - 2 * pad)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _runs(tf.paragraphs[0], el["desc"], el["size"])


def _badge(slide, label, color):
    w = 7.2 * len(label) + 20
    s = _rect(slide, theme.SLIDE_W - 43 - w, 30, w, 20, color, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = s.text_frame
    tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _runs(p, label, 10, color="white", bold=True)


def _footer(slide, deck, n):
    m = deck.meta
    _, tf = _box(slide, 43, theme.FOOTER_Y, 700, 14)
    _runs(tf.paragraphs[0], f"{m.get('course', '')} · Session {m.get('session', '')}: {m.get('title', '')}",
          theme.FOOTER_SIZE, color="muted")
    _, tf = _box(slide, theme.SLIDE_W - 43 - 60, theme.FOOTER_Y, 60, 14)
    tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
    _runs(tf.paragraphs[0], str(n), theme.FOOTER_SIZE, color="muted")


def _title_slide(slide, deck):
    m = deck.meta
    _rect(slide, 0, 0, theme.SLIDE_W, theme.SLIDE_H, "primary")
    _rect(slide, 60, 250, 80, 4, "accent")
    _, tf = _box(slide, 60, 120, 840, 30)
    _runs(tf.paragraphs[0], m.get("course", ""), 20, color="white")
    _, tf = _box(slide, 60, 160, 840, 80)
    tf.vertical_anchor = MSO_ANCHOR.BOTTOM
    _runs(tf.paragraphs[0], f"Session {m.get('session', '')}: {m.get('title', '')}", 40, color="white", bold=True)
    lines = [m.get("subtitle", ""), m.get("program", ""), m.get("lecturer", "")]
    _, tf = _box(slide, 60, 270, 840, 120)
    for i, ln in enumerate(l for l in lines if l):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        _runs(p, ln, 18 if i == 0 else 14, color="white")


def _section_slide(slide, sl):
    _rect(slide, 0, 0, theme.SLIDE_W, theme.SLIDE_H, "primary")
    _rect(slide, 60, 300, 80, 4, "accent")
    _, tf = _box(slide, 60, 180, 840, 110)
    tf.vertical_anchor = MSO_ANCHOR.BOTTOM
    _runs(tf.paragraphs[0], sl.title, 36, color="white", bold=True)
    for b in sl.blocks:
        if b["type"] == "para":
            _, tf = _box(slide, 60, 320, 840, 80)
            _runs(tf.paragraphs[0], b["text"], 18, color="white")


def render(deck, out_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(theme.SLIDE_W), Pt(theme.SLIDE_H)
    blank = prs.slide_layouts[6]
    for n, sl in enumerate(deck.slides, start=1):
        slide = prs.slides.add_slide(blank)
        if "title" in sl.classes:
            _title_slide(slide, deck)
        elif "section" in sl.classes:
            _section_slide(slide, sl)
        else:
            _, tf = _box(slide, *theme.TITLE)
            tf.vertical_anchor = MSO_ANCHOR.BOTTOM
            _runs(tf.paragraphs[0], sl.title, theme.TITLE_SIZE, color="primary", bold=True)
            _rect(slide, theme.TITLE[0], theme.RULE_Y, 60, 3, "accent")
            if "optional" in sl.classes:
                _badge(slide, "OPTIONAL · SELF-STUDY", "optional")
            elif "exercise" in sl.classes:
                _badge(slide, "EXERCISE", "accent")
            refs = "references" in sl.classes
            for el in sl.layout:
                k = el["kind"]
                if k == "text":
                    _text(slide, el, refs=refs)
                elif k == "table":
                    _table(slide, el)
                elif k == "code":
                    _code(slide, el)
                elif k == "card":
                    _rect(slide, el["x"], el["y"], el["w"], el["h"], "light", "border", MSO_SHAPE.ROUNDED_RECTANGLE)
                elif k == "callout":
                    _rect(slide, el["x"], el["y"], el["w"], el["h"], "light")
                    _rect(slide, el["x"], el["y"], theme.CALLOUT_BAR, el["h"], "accent")
                elif k == "layer":
                    _layer(slide, el)
                elif k == "image":
                    slide.shapes.add_picture(el["src"], Pt(el["x"]), Pt(el["y"]), Pt(el["w"]), Pt(el["h"]))
            _footer(slide, deck, n)
        if sl.notes:
            slide.notes_slide.notes_text_frame.text = sl.notes
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
